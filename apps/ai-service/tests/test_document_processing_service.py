"""Unit tests for DocumentProcessingService orchestrator, polymorphic routing, and defenses."""

import os
import tempfile
from unittest.mock import AsyncMock, MagicMock, patch

import docx
import pymupdf
import pytest
import pytest_asyncio
from bson import ObjectId
from qdrant_client import AsyncQdrantClient
from src.core.exceptions import (
    DocumentNotFoundError,
    PathTraversalSecurityError,
)
from src.infrastructure.ai.embedding_provider import IEmbeddingProvider
from src.infrastructure.database.qdrant_client import init_qdrant_schema
from src.infrastructure.database.vector_repository import (
    VectorPoint,
    VectorRepository,
)
from src.services.document_processing_service import (
    PARSER_REGISTRY,
    DocumentProcessingService,
    resolve_safe_filepath,
)


@pytest.fixture
def mock_embedding_service():
    """Provides a mocked embedding service returning deterministic 384-d vectors."""
    mock_service = MagicMock(spec=IEmbeddingProvider)
    mock_service.model_name = "bge-small-en-v1.5"
    mock_service.dimension = 384
    mock_service.device = "cpu"
    mock_service.embed_batch.side_effect = lambda texts: [[0.05] * 384 for _ in texts]
    mock_service.embed_text.return_value = [0.05] * 384
    return mock_service


@pytest.fixture
def mock_vector_repo():
    """Provides a mocked VectorRepository recording upserted points."""
    repo = MagicMock(spec=VectorRepository)
    repo.upsert_vectors = AsyncMock(side_effect=lambda points: len(points))
    return repo


@pytest_asyncio.fixture
async def memory_qdrant_repo():
    """Provisions a real VectorRepository hooked to an isolated in-memory Qdrant client."""
    client = AsyncQdrantClient(":memory:")
    await init_qdrant_schema(client=client)
    repo = VectorRepository(client=client, collection_name="academic_documents")
    yield repo
    await client.close()


def test_resolve_safe_filepath_valid():
    """Verify safe path resolution strips prefixes and returns target inside upload dir."""
    upload_dir = os.path.abspath("test_uploads")
    resolved = resolve_safe_filepath("/uploads/my-paper.pdf", upload_dir=upload_dir)
    assert resolved == os.path.join(upload_dir, "my-paper.pdf")

    resolved_bare = resolve_safe_filepath("my-paper.pdf", upload_dir=upload_dir)
    assert resolved_bare == os.path.join(upload_dir, "my-paper.pdf")


def test_resolve_safe_filepath_traversal_attempts():
    """Verify various path traversal attempts are detected and rejected."""
    upload_dir = os.path.abspath("test_uploads")

    with pytest.raises(PathTraversalSecurityError):
        resolve_safe_filepath("../../etc/passwd", upload_dir=upload_dir)

    with pytest.raises(PathTraversalSecurityError):
        resolve_safe_filepath("uploads/../../secret.txt", upload_dir=upload_dir)

    with pytest.raises(PathTraversalSecurityError):
        resolve_safe_filepath("", upload_dir=upload_dir)


@pytest.mark.asyncio
async def test_process_document_invalid_object_id():
    """Verify invalid document ID raises DocumentNotFoundError."""
    service = DocumentProcessingService(db=MagicMock())
    with pytest.raises(DocumentNotFoundError):
        await service.process_document("invalid-id-format", "uploads/test.pdf")


@pytest.mark.asyncio
async def test_process_document_not_found_in_db():
    """Verify DocumentNotFoundError is raised when record does not exist in DB."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    mock_collection.find_one_and_update = AsyncMock(return_value=None)

    service = DocumentProcessingService(db=mock_db)
    valid_id = str(ObjectId())

    with pytest.raises(DocumentNotFoundError):
        await service.process_document(valid_id, "uploads/test.pdf")


@pytest.mark.asyncio
async def test_process_pdf_document_success(mock_embedding_service, mock_vector_repo):
    """Verify PDF processing executes full pipeline, chunking, embedding, upsert, and COMPLETED status."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    user_id = str(ObjectId())
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "userId": user_id,
            "mimeType": "application/pdf",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    with tempfile.TemporaryDirectory() as tmp_dir:
        # Create a real small PDF
        pdf_path = os.path.join(tmp_dir, "academic_test.pdf")
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text(
            (50, 100), "This is an academic research document on RAG.", fontsize=11
        )
        doc.save(pdf_path)
        doc.close()

        service = DocumentProcessingService(
            db=mock_db,
            embedding_service=mock_embedding_service,
            vector_repository=mock_vector_repo,
        )

        with patch(
            "src.services.document_processing_service.settings"
        ) as mock_settings:
            mock_settings.UPLOAD_DIR = tmp_dir
            await service.process_document(doc_id, "/uploads/academic_test.pdf")

    # Verify find_one_and_update was called to set PARSING
    mock_collection.find_one_and_update.assert_awaited_once()

    # Verify update_one was called: first to EMBEDDING, then to COMPLETED
    assert mock_collection.update_one.await_count == 2
    first_call = mock_collection.update_one.call_args_list[0][0]
    second_call = mock_collection.update_one.call_args_list[1][0]

    assert first_call[0] == {"_id": ObjectId(doc_id)}
    assert first_call[1]["$set"]["status"] == "EMBEDDING"
    assert first_call[1]["$set"]["metadata.pageCount"] == 1
    assert first_call[1]["$set"]["metadata.wordCount"] > 0
    assert first_call[1]["$set"]["metadata.chunkCount"] >= 1

    assert second_call[0] == {"_id": ObjectId(doc_id)}
    assert second_call[1]["$set"]["status"] == "COMPLETED"

    # Verify embeddings and vector repository upsert were executed
    mock_embedding_service.embed_batch.assert_called_once()
    mock_vector_repo.upsert_vectors.assert_awaited_once()
    upserted_points = mock_vector_repo.upsert_vectors.call_args[0][0]
    assert len(upserted_points) >= 1
    assert isinstance(upserted_points[0], VectorPoint)
    assert upserted_points[0].document_id == doc_id
    assert upserted_points[0].user_id == user_id


@pytest.mark.asyncio
async def test_process_docx_document_success(mock_embedding_service, mock_vector_repo):
    """Verify DOCX processing routes to DocxParser, populates chunkCount, unsets pageCount, and completes."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    user_id = str(ObjectId())
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "userId": user_id,
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    with tempfile.TemporaryDirectory() as tmp_dir:
        # Create a real small DOCX
        docx_path = os.path.join(tmp_dir, "thesis_draft.docx")
        doc = docx.Document()
        doc.add_paragraph("Chapter 1: Neural Networks in Modern Education")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Metric"
        table.rows[0].cells[1].text = "Value"
        doc.save(docx_path)

        service = DocumentProcessingService(
            db=mock_db,
            embedding_service=mock_embedding_service,
            vector_repository=mock_vector_repo,
        )

        with patch(
            "src.services.document_processing_service.settings"
        ) as mock_settings:
            mock_settings.UPLOAD_DIR = tmp_dir
            await service.process_document(doc_id, "/uploads/thesis_draft.docx")

    assert mock_collection.update_one.await_count == 2
    first_call = mock_collection.update_one.call_args_list[0][0]
    second_call = mock_collection.update_one.call_args_list[1][0]

    assert first_call[1]["$set"]["status"] == "EMBEDDING"
    assert first_call[1]["$set"]["metadata.wordCount"] > 5
    assert first_call[1]["$set"]["metadata.chunkCount"] >= 1
    assert "metadata.pageCount" not in first_call[1]["$set"]
    assert first_call[1]["$unset"]["metadata.pageCount"] == ""

    assert second_call[1]["$set"]["status"] == "COMPLETED"
    mock_vector_repo.upsert_vectors.assert_awaited_once()


@pytest.mark.asyncio
async def test_process_txt_document_success(mock_embedding_service, mock_vector_repo):
    """Verify TXT processing routes to TxtParser, populates chunkCount, and reaches COMPLETED status."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    user_id = str(ObjectId())
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "userId": user_id,
            "mimeType": "text/plain",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    with tempfile.TemporaryDirectory() as tmp_dir:
        txt_path = os.path.join(tmp_dir, "notes.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Dataset notes and LaTeX equations for analysis.")

        service = DocumentProcessingService(
            db=mock_db,
            embedding_service=mock_embedding_service,
            vector_repository=mock_vector_repo,
        )

        with patch(
            "src.services.document_processing_service.settings"
        ) as mock_settings:
            mock_settings.UPLOAD_DIR = tmp_dir
            await service.process_document(doc_id, "/uploads/notes.txt")

    assert mock_collection.update_one.await_count == 2
    first_call = mock_collection.update_one.call_args_list[0][0]
    second_call = mock_collection.update_one.call_args_list[1][0]

    assert first_call[1]["$set"]["status"] == "EMBEDDING"
    assert first_call[1]["$set"]["metadata.wordCount"] > 5
    assert first_call[1]["$set"]["metadata.chunkCount"] >= 1
    assert "metadata.pageCount" not in first_call[1]["$set"]
    assert first_call[1]["$unset"]["metadata.pageCount"] == ""

    assert second_call[1]["$set"]["status"] == "COMPLETED"


@pytest.mark.asyncio
async def test_process_txt_with_charset_parameter(
    mock_embedding_service, mock_vector_repo
):
    """Verify MIME types with parameters (e.g. text/plain; charset=utf-8) route successfully."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "mimeType": "text/plain; charset=utf-8",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    with tempfile.TemporaryDirectory() as tmp_dir:
        txt_path = os.path.join(tmp_dir, "charset_notes.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Valid UTF-8 plain text with charset parameter in MIME header.")

        service = DocumentProcessingService(
            db=mock_db,
            embedding_service=mock_embedding_service,
            vector_repository=mock_vector_repo,
        )

        with patch(
            "src.services.document_processing_service.settings"
        ) as mock_settings:
            mock_settings.UPLOAD_DIR = tmp_dir
            await service.process_document(doc_id, "/uploads/charset_notes.txt")

    assert mock_collection.update_one.await_count == 2
    assert (
        mock_collection.update_one.call_args_list[1][0][1]["$set"]["status"]
        == "COMPLETED"
    )


@pytest.mark.asyncio
async def test_process_document_unsupported_mime_type():
    """Verify unsupported MIME type transitions document status to FAILED."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "mimeType": "image/png",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    service = DocumentProcessingService(db=mock_db)

    with patch("src.services.document_processing_service.settings") as mock_settings:
        mock_settings.UPLOAD_DIR = "test_uploads"
        await service.process_document(doc_id, "/uploads/unsupported.png")

    mock_collection.update_one.assert_awaited_once()
    call_args = mock_collection.update_one.call_args[0]
    assert call_args[1]["$set"]["status"] == "FAILED"
    assert "unsupported" in call_args[1]["$set"]["errorMessage"].lower()


@pytest.mark.asyncio
async def test_process_document_failure_transitions_to_failed():
    """Verify extraction failure (missing file) transitions document status to FAILED."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "mimeType": "application/pdf",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    service = DocumentProcessingService(db=mock_db)

    # Point to a non-existent file to simulate failure
    with patch("src.services.document_processing_service.settings") as mock_settings:
        mock_settings.UPLOAD_DIR = "/fake/uploads"
        await service.process_document(doc_id, "/uploads/non_existent.pdf")

    mock_collection.update_one.assert_awaited_once()
    call_args = mock_collection.update_one.call_args[0]
    assert call_args[1]["$set"]["status"] == "FAILED"
    assert "errorMessage" in call_args[1]["$set"]
    assert "not found" in call_args[1]["$set"]["errorMessage"].lower()


def test_parser_registry_contains_supported_mimes():
    """Verify PARSER_REGISTRY contains PDF, DOCX, and TXT strategies."""
    assert "application/pdf" in PARSER_REGISTRY
    assert (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        in PARSER_REGISTRY
    )
    assert "text/plain" in PARSER_REGISTRY


@pytest.mark.asyncio
async def test_process_document_populates_extracted_metadata_and_augmented_embeddings(
    mock_embedding_service, mock_vector_repo
):
    """Verify academic document processing extracts metadata, uses augmented embeddings with title, and completes."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    user_id = "user_scholar_42"
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "userId": user_id,
            "mimeType": "text/plain",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    with tempfile.TemporaryDirectory() as tmp_dir:
        doc_path = os.path.join(tmp_dir, "academic_paper.txt")
        with open(doc_path, "w", encoding="utf-8") as f:
            f.write(
                "Transformer Models for Academic Reasoning\n\n"
                "Authors: Alan Turing, Ada Lovelace\n\n"
                "Abstract\n\n"
                "We demonstrate self-attention mechanisms applied to scientific discovery.\n\n"
                "Keywords: transformers, reasoning, ai\n\n"
                "1. Introduction\n\n"
                "Foundation models continue to evolve."
            )

        service = DocumentProcessingService(
            db=mock_db,
            embedding_service=mock_embedding_service,
            vector_repository=mock_vector_repo,
        )

        with patch(
            "src.services.document_processing_service.settings"
        ) as mock_settings:
            mock_settings.UPLOAD_DIR = tmp_dir
            await service.process_document(doc_id, "/uploads/academic_paper.txt")

    assert mock_collection.update_one.await_count == 2
    first_call = mock_collection.update_one.call_args_list[0][0]
    set_fields = first_call[1]["$set"]
    assert set_fields["status"] == "EMBEDDING"
    assert (
        set_fields["metadata.title_extracted"]
        == "Transformer Models for Academic Reasoning"
    )
    assert set_fields["metadata.authors"] == ["Alan Turing", "Ada Lovelace"]
    assert (
        set_fields["metadata.abstract"]
        == "We demonstrate self-attention mechanisms applied to scientific discovery."
    )
    assert set_fields["metadata.keywords"] == ["transformers", "reasoning", "ai"]
    assert set_fields["metadata.chunkCount"] >= 1

    # Verify Augmented Chunking: title prepended in embedding input
    mock_embedding_service.embed_batch.assert_called_once()
    embedded_texts = mock_embedding_service.embed_batch.call_args[0][0]
    assert any(
        "Title: Transformer Models for Academic Reasoning" in t for t in embedded_texts
    )

    # Verify Vector Repository received VectorPoint with title metadata
    mock_vector_repo.upsert_vectors.assert_awaited_once()
    upserted_points = mock_vector_repo.upsert_vectors.call_args[0][0]
    assert (
        upserted_points[0].metadata["title"]
        == "Transformer Models for Academic Reasoning"
    )
    assert upserted_points[0].user_id == user_id

    # Verify final status transition
    second_call = mock_collection.update_one.call_args_list[1][0]
    assert second_call[1]["$set"]["status"] == "COMPLETED"


@pytest.mark.asyncio
async def test_process_document_unstructured_txt_safe_fallback(
    mock_embedding_service, mock_vector_repo
):
    """Verify unstructured TXT sets abstract and keywords to None without throwing and completes."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "mimeType": "text/plain",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    with tempfile.TemporaryDirectory() as tmp_dir:
        doc_path = os.path.join(tmp_dir, "unstructured_notes.txt")
        with open(doc_path, "w", encoding="utf-8") as f:
            f.write(
                "Weekly engineering sync notes.\nNo abstract or formal keywords here."
            )

        service = DocumentProcessingService(
            db=mock_db,
            embedding_service=mock_embedding_service,
            vector_repository=mock_vector_repo,
        )

        with patch(
            "src.services.document_processing_service.settings"
        ) as mock_settings:
            mock_settings.UPLOAD_DIR = tmp_dir
            await service.process_document(doc_id, "/uploads/unstructured_notes.txt")

    assert mock_collection.update_one.await_count == 2
    first_call = mock_collection.update_one.call_args_list[0][0]
    set_fields = first_call[1]["$set"]
    assert set_fields["status"] == "EMBEDDING"
    assert set_fields["metadata.title_extracted"] == "Weekly engineering sync notes."
    assert set_fields["metadata.abstract"] is None
    assert set_fields["metadata.keywords"] is None

    second_call = mock_collection.update_one.call_args_list[1][0]
    assert second_call[1]["$set"]["status"] == "COMPLETED"


@pytest.mark.asyncio
async def test_process_document_with_version_id_positional_update(
    mock_embedding_service, mock_vector_repo
):
    """Verify processing with version_id updates matching nested version via positional operator to COMPLETED."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    user_id = str(ObjectId())
    version_id = "test-ver-uuid-123"
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "userId": user_id,
            "mimeType": "text/plain",
            "status": "PARSING",
            "versions": [
                {
                    "versionId": version_id,
                    "mimeType": "text/plain",
                    "status": "PARSING",
                }
            ],
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    with tempfile.TemporaryDirectory() as tmp_dir:
        doc_path = os.path.join(tmp_dir, "version2_paper.txt")
        with open(doc_path, "w", encoding="utf-8") as f:
            f.write(
                "Deep Reinforcement Learning Revision 2\n\nAbstract: Improved stability."
            )

        service = DocumentProcessingService(
            db=mock_db,
            embedding_service=mock_embedding_service,
            vector_repository=mock_vector_repo,
        )

        with patch(
            "src.services.document_processing_service.settings"
        ) as mock_settings:
            mock_settings.UPLOAD_DIR = tmp_dir
            await service.process_document(
                doc_id, "/uploads/version2_paper.txt", version_id=version_id
            )

    # Verify find_one_and_update filtered with versions.versionId
    find_args = mock_collection.find_one_and_update.call_args[0]
    assert find_args[0] == {"_id": ObjectId(doc_id), "versions.versionId": version_id}
    assert find_args[1]["$set"]["versions.$.status"] == "PARSING"

    # Verify update_one targeted positional versions.$.metadata and versions.$.status
    assert mock_collection.update_one.await_count == 2
    first_call = mock_collection.update_one.call_args_list[0][0]
    assert first_call[0] == {"_id": ObjectId(doc_id), "versions.versionId": version_id}
    set_fields = first_call[1]["$set"]
    assert set_fields["status"] == "EMBEDDING"
    assert set_fields["versions.$.status"] == "EMBEDDING"
    assert "versions.$.metadata" in set_fields
    assert set_fields["versions.$.metadata"]["wordCount"] > 0
    assert set_fields["versions.$.metadata"]["chunkCount"] >= 1
    assert (
        set_fields["versions.$.metadata"]["title_extracted"]
        == "Deep Reinforcement Learning Revision 2"
    )

    second_call = mock_collection.update_one.call_args_list[1][0]
    assert second_call[0] == {
        "_id": ObjectId(doc_id),
        "versions.versionId": version_id,
    }
    assert second_call[1]["$set"]["status"] == "COMPLETED"
    assert second_call[1]["$set"]["versions.$.status"] == "COMPLETED"

    # Verify vector point has version_id
    upserted_points = mock_vector_repo.upsert_vectors.call_args[0][0]
    assert upserted_points[0].version_id == version_id


@pytest.mark.asyncio
async def test_process_document_end_to_end_with_memory_qdrant(
    memory_qdrant_repo, mock_embedding_service
):
    """End-to-end integration test verifying that DocumentProcessingService indexes into real Qdrant collection."""
    mock_db = MagicMock()
    mock_collection = MagicMock()
    mock_db.documents = mock_collection

    doc_id = str(ObjectId())
    user_id = "user_e2e_researcher"
    mock_collection.find_one_and_update = AsyncMock(
        return_value={
            "_id": ObjectId(doc_id),
            "userId": user_id,
            "mimeType": "text/plain",
            "status": "PARSING",
        }
    )
    mock_collection.update_one = AsyncMock(return_value=MagicMock(matched_count=1))

    with tempfile.TemporaryDirectory() as tmp_dir:
        doc_path = os.path.join(tmp_dir, "e2e_paper.txt")
        with open(doc_path, "w", encoding="utf-8") as f:
            f.write(
                "Autonomous Multi-Agent Architecture for Scientific Discovery\n\n"
                "Abstract: We demonstrate recursive token chunking and dense vector indexing.\n\n"
                "Section 1: The foundation models interface with vector databases seamlessly."
            )

        service = DocumentProcessingService(
            db=mock_db,
            embedding_service=mock_embedding_service,
            vector_repository=memory_qdrant_repo,
        )

        with patch(
            "src.services.document_processing_service.settings"
        ) as mock_settings:
            mock_settings.UPLOAD_DIR = tmp_dir
            await service.process_document(doc_id, "/uploads/e2e_paper.txt")

    # Verify status transitions
    assert mock_collection.update_one.await_count == 2
    final_status = mock_collection.update_one.call_args_list[1][0][1]["$set"]["status"]
    assert final_status == "COMPLETED"

    # Perform similarity search on in-memory Qdrant repository to verify data persistence
    results = await memory_qdrant_repo.search_similarity(
        query_vector=[0.05] * 384,
        user_id=user_id,
        document_id=doc_id,
        limit=5,
    )
    assert len(results) >= 1
    assert results[0].document_id == doc_id
    assert results[0].user_id == user_id
    assert "Autonomous Multi-Agent" in results[0].text
    assert (
        results[0].metadata["title"]
        == "Autonomous Multi-Agent Architecture for Scientific Discovery"
    )
