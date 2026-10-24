"""High-fidelity Microsoft Word (.docx) document extractor including structural table parsing."""

import logging
import os
import zipfile

import docx
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.core.exceptions import (
    CorruptedDOCXError,
    EmptyDocumentError,
    FileNotFoundDOCXError,
)
from src.pipelines.extraction.parser_interface import BaseParser

logger = logging.getLogger(__name__)


def extract_text_from_docx(filepath: str) -> tuple[str, int | None, int]:
    """Extracts raw text content from a .docx file adhering to sequential document body flow.

    Extracts text from both standard paragraphs and embedded tables, deduplicating
    horizontally merged table cells and preserving natural layout ordering.

    Args:
        filepath: Absolute or validated local path to the .docx document.

    Returns:
        Tuple of (raw_text, page_count=None, approximate_word_count).

    Raises:
        FileNotFoundDOCXError: When the target file does not exist.
        CorruptedDOCXError: When the file is not a valid zip archive or has corrupted OpenXML.
        EmptyDocumentError: When the file contains no readable text content.
    """
    if not os.path.isfile(filepath):
        raise FileNotFoundDOCXError(f"DOCX file not found at: {filepath}")

    try:
        doc = docx.Document(filepath)
    except (
        PackageNotFoundError,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        Exception,
    ) as exc:
        logger.error(f"Failed to open DOCX archive {filepath}: {exc}")
        raise CorruptedDOCXError(
            f"Invalid or corrupted DOCX document archive: {os.path.basename(filepath)}"
        ) from exc

    extracted_blocks: list[str] = []

    try:
        # Iterate over document body elements in natural order (paragraphs and tables)
        for child in doc.element.body.iterchildren():
            # Paragraph element (w:p)
            if child.tag.endswith("p"):
                paragraph = Paragraph(child, doc)
                text = paragraph.text.strip()
                if text:
                    extracted_blocks.append(text)
            # Table element (w:tbl)
            elif child.tag.endswith("tbl"):
                table = Table(child, doc)
                for row in table.rows:
                    seen_tcs = set()
                    row_cells: list[str] = []
                    for cell in row.cells:
                        # Deduplicate merged cell references in same row
                        if cell._tc not in seen_tcs:
                            seen_tcs.add(cell._tc)
                            cell_text = cell.text.strip()
                            if cell_text:
                                row_cells.append(cell_text)
                    if row_cells:
                        extracted_blocks.append(" | ".join(row_cells))
    except Exception as body_err:
        logger.warning(
            f"Sequential body parsing encountered an issue in {filepath}: {body_err}. Falling back to standard paragraph/table iteration."
        )
        # Fallback to standard doc.paragraphs + doc.tables if body iterchildren encounters an unexpected XML construct
        extracted_blocks = []
        for p in doc.paragraphs:
            if p.text.strip():
                extracted_blocks.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                seen_tcs = set()
                row_cells = []
                for cell in row.cells:
                    if cell._tc not in seen_tcs:
                        seen_tcs.add(cell._tc)
                        if cell.text.strip():
                            row_cells.append(cell.text.strip())
                if row_cells:
                    extracted_blocks.append(" | ".join(row_cells))

    if not extracted_blocks:
        raise EmptyDocumentError(
            f"DOCX document contains no extractable text: {os.path.basename(filepath)}"
        )

    full_text = "\n\n".join(extracted_blocks).strip()
    word_count = len(full_text.split())

    # DOCX is a flow format without fixed pages; page count is inherently dynamic/None
    return full_text, None, word_count


class DocxParser(BaseParser):
    """Microsoft Word (.docx) document parser strategy."""

    def extract(self, filepath: str) -> tuple[str, int | None, int]:
        """Extracts text content and metadata from a Word document."""
        return extract_text_from_docx(filepath)
